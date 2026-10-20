"""Component driver for the pinned MiMo DOCX core-metrics rule.

Requires python-docx in the execution image. The metadata is supplied as a
read-only source path; this driver never downloads or installs dependencies.
"""

import argparse
import json
from pathlib import Path
import sys

from docx import Document


MEMO = Path("03_Reporting/Reimbursement_Committee_Reconciliation_Memo.docx")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--metadata", required=True)
    args = parser.parse_args()
    payload = json.load(sys.stdin)
    data = payload["case"]["input"]
    table_data = data["table"]
    if not isinstance(table_data, list) or not table_data:
        raise ValueError("table must be a nonempty list of [label, value] pairs")
    workspace = Path(payload["workspace"])
    path = workspace / MEMO
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading(str(data.get("title", "Committee reconciliation metrics")), level=1)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    for pair in table_data:
        if not isinstance(pair, list) or len(pair) != 2 or not all(isinstance(x, str) for x in pair):
            raise ValueError("each table row must be [text, text]")
        cells = table.add_row().cells
        cells[0].text, cells[1].text = pair
    document.save(path)
    metadata = json.loads(Path(args.metadata).read_text())
    namespace = {}
    exec(compile(metadata["check_code"], args.metadata + ":check_code", "exec"), namespace)
    result = namespace["check_memo_core_metrics"](workspace)
    print(json.dumps({"reward": float(result["score"]),
                      "details": {"rule": "check_memo_core_metrics", "result": result}}, allow_nan=False))


if __name__ == "__main__":
    main()
