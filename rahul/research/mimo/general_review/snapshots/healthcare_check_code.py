import hashlib
import re
from pathlib import Path

from docx import Document


MEMO = Path("03_Reporting/Reimbursement_Committee_Reconciliation_Memo.docx")
SOURCE_HASHES = {
    "02_Validation_Workpapers/Transfer_Payment_Reconciliation.xlsx": "9f7c07a1b151997a8f654226abde299ceed62e993c6afd0ba7aa40cb66106b26",
    "03_Reporting/September_2025_DRG_Transfer_Validation_Report.docx": "dca8f04bb35548de9bff073c8b75835851869ff04068b01048e9a53272a50807",
    "03_Reporting/Revenue_Cycle_Transfer_Risk_Memo.docx": "74dca9610f88db0229233c5d49885987cbf9be14411750d7c091629529eaa1b6",
}


def _memo_path(ws):
    p = ws / MEMO
    if p.exists():
        return p
    # Accept equivalent committee memo filenames in the reporting area.
    return locate_deliverable(ws, "Reconciliation", "Memo", exts=(".docx",))


def _text(ws):
    p = _memo_path(ws)
    if not p or not p.exists():
        return ""
    doc = Document(p)
    parts = [x.text for x in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _money(text, value):
    dollars = f"{abs(value):,.2f}"
    signed = [dollars, f"${dollars}"]
    if value < 0:
        signed += [f"-${dollars}", f"({dollars})", f"({dollars.replace(',', '')})"]
    return any(token in text for token in signed)


def _has_number(text, number):
    return re.search(rf"(?<!\d){number}(?!\d)", text) is not None


def check_memo_core_metrics(ws):
    text = _text(ws)
    checks = [
        ("24", _has_number(text, 24)),
        ("9", _has_number(text, 9)),
        ("7", _has_number(text, 7)),
        ("184672.54", _money(text, 184672.54)),
        ("162826.17", _money(text, 162826.17)),
        ("21846.37 decrease", _money(text, -21846.37) and bool(re.search(r"decrease|reduction|lower", text, re.I))),
    ]
    missing = [name for name, ok in checks if not ok]
    return {"score": int(not missing), "detail": "missing: " + ", ".join(missing) if missing else "core metrics present"}


def check_memo_cross_source_control(ws):
    text = _text(ws)
    checks = [
        ("run id", "RR-FY25Q1-TRANSFER-044" in text),
        ("finalized", bool(re.search(r"FINALIZED|finalized", text))),
        ("Epic/report 24", bool(re.search(r"Epic.{0,100}24|24.{0,100}Epic|report.{0,100}24|24.{0,100}report", text, re.I | re.S))),
        ("StrataJazz 9", bool(re.search(r"StrataJazz.{0,100}9|9.{0,100}StrataJazz", text, re.I | re.S))),
        ("difference 15", bool(re.search(r"15[- ]case|difference.{0,40}15|15.{0,40}difference|scope difference", text, re.I | re.S)) or (bool(re.search(r"(?:24|twenty[- ]four).{0,120}(?:9|nine)|(?:9|nine).{0,120}(?:24|twenty[- ]four)", text, re.I | re.S)) and bool(re.search(r"subset|non[- ]transfer|does not include|not replace|different denominator|payment subset", text, re.I | re.S)))),
        ("seven lines", bool(re.search(r"7.{0,40}(correction|corrected).{0,40}(line|claim)|(?:correction|corrected).{0,40}7.{0,40}(line|claim)", text, re.I | re.S))),
        ("financial match", _money(text, 162826.17) and _money(text, -21846.37)),
    ]
    missing = [name for name, ok in checks if not ok]
    return {"score": int(not missing), "detail": "missing: " + ", ".join(missing) if missing else "cross-source control disclosed"}


def check_memo_action_conclusion(ws):
    text = _text(ws)
    checks = [
        ("seven adjustment bills", bool(re.search(r"7|seven", text, re.I)) and bool(re.search(r"adjustment bill", text, re.I))),
        ("FL17 01", bool(re.search(r"FL17.{0,20}01|01.{0,20}FL17", text, re.I))),
        ("FL17 02", bool(re.search(r"FL17.{0,20}02|02.{0,20}FL17", text, re.I))),
        ("preserve audit trail", bool(re.search(r"preserv.{0,30}(original|audit)|original.{0,30}(audit|claim) history", text, re.I))),
        ("30-case October control", bool(re.search(r"30[- ]case|sample.{0,50}30|30.{0,80}(?:Medicare|FFS).{0,80}(?:inpatient|discharge)", text, re.I | re.S)) and bool(re.search(r"October 2025", text, re.I))),
    ]
    missing = [name for name, ok in checks if not ok]
    return {"score": int(not missing), "detail": "missing: " + ", ".join(missing) if missing else "action and cadence present"}


def check_source_integrity(ws):
    bad = []
    for rel, expected in SOURCE_HASHES.items():
        p = ws / rel
        actual = hashlib.sha256(p.read_bytes()).hexdigest() if p.exists() else "missing"
        if actual != expected:
            bad.append(rel)
    return {"score": int(not bad), "detail": "changed: " + ", ".join(bad) if bad else "protected sources unchanged"}
