#!/usr/bin/env python3
"""Paired DOCX component probe using the published General rule unchanged.

Run through Slurm on amd-tw-verification with python-docx installed in an
isolated dependency environment. This does not call any LLM or full runtime.
"""
import hashlib
import importlib.metadata
import json
import os
from pathlib import Path
import platform

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
TASK_ID = "s3k_1401_healthcare_ops_en_t4_rl_007"
SOURCE = ROOT / "research/mimo/general_actions/snapshots" / TASK_ID
OUT = ROOT / "research/mimo/probes/money-format"
MEMO = Path("03_Reporting/Reimbursement_Committee_Reconciliation_Memo.docx")


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    if os.environ.get("SLURM_JOB_PARTITION") != "amd-tw-verification":
        raise SystemExit("Run through Slurm in amd-tw-verification")
    meta = json.loads((SOURCE / "verifier_meta.json").read_text())
    namespace = {}
    # Reviewed code defines DOCX readers and fixed checks, with no top-level
    # task execution. The selected check reads only the generated memo.
    exec(compile(meta["check_code"], str(SOURCE / "verifier_meta.json") + ":check_code", "exec"), namespace)
    job = os.environ["SLURM_JOB_ID"]
    result_dir = OUT / job
    result_dir.mkdir(parents=True, exist_ok=False)
    observations = []
    values = [
        ("Reviewed population", "24"),
        ("Transfer-rule population", "9"),
        ("Affected claims", "7"),
        ("Current billed amount", 184672.54),
        ("Corrected transfer amount", 162826.17),
        ("Estimated reimbursement decrease", -21846.37),
    ]
    for variant, grouping in [("commas", True), ("no_commas", False)]:
        workspace = result_dir / variant
        path = workspace / MEMO
        path.parent.mkdir(parents=True)
        document = Document()
        document.add_heading("Committee reconciliation metrics", level=1)
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Metric"
        table.rows[0].cells[1].text = "Value"
        rendered = []
        for label, value in values:
            if isinstance(value, float) and value < 0:
                # Hold the negative amount fixed to isolate the check's
                # treatment of the two positive monetary values.
                text = f"(${abs(value):,.2f})"
            elif isinstance(value, float):
                text = f"${value:,.2f}" if grouping else f"${value:.2f}"
            else:
                text = value
            cells = table.add_row().cells
            cells[0].text, cells[1].text = label, text
            rendered.append([label, text])
        document.save(path)
        result = namespace["check_memo_core_metrics"](workspace)
        observations.append({
            "variant": variant,
            "table": rendered,
            "rule_result": result,
            "artifact": str(path.relative_to(ROOT)),
            "artifact_sha256": digest(path),
        })
    assert observations[0]["rule_result"]["score"] == 1, observations
    assert observations[1]["rule_result"]["score"] == 0, observations
    # Verify that paired tables have identical numeric meanings independently
    # of the rule's string matcher.
    from decimal import Decimal
    for left, right in zip(observations[0]["table"], observations[1]["table"]):
        assert left[0] == right[0]
        def numeric(text):
            cleaned = text.replace("$", "").replace(",", "")
            return -Decimal(cleaned[1:-1]) if cleaned.startswith("(") else Decimal(cleaned)
        assert numeric(left[1]) == numeric(right[1])
    report = {
        "id": "G4_money_grouping_false_rejection",
        "task_id": TASK_ID,
        "dataset_revision": "e8e81e4945270b6212204bad99de1e598f7fb077",
        "status": "confirmed_component_counterexample",
        "job_id": job,
        "partition": os.environ["SLURM_JOB_PARTITION"],
        "source_hashes": {name: digest(SOURCE / name) for name in ["instruction.md", "verifier_meta.json"]},
        "dependencies": {"python": platform.python_version(), "python-docx": importlib.metadata.version("python-docx"), "lxml": importlib.metadata.version("lxml")},
        "observations": observations,
        "interpretation": "The unmodified core-metrics rule accepts grouped positive monetary values and rejects numerically equivalent ungrouped values. All other text and the negative amount are held fixed. The prompt does not specify a digit-grouping convention.",
        "limitations": [
            "Component-level rule execution using paired synthetic DOCX metric tables; neither is a complete task submission.",
            "No LLM criterion, task container, source-integrity rule or trainer transformation was run.",
            "This demonstrates a numeric-format counterexample for the core-metrics check, not a measured total reward or training effect.",
        ],
    }
    output = result_dir / "result.json"
    output.write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps({"output": str(output), "scores": [x["rule_result"]["score"] for x in observations]}))


if __name__ == "__main__":
    main()
